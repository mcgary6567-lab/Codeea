"""Interactive reports (docs/INTERACTIVE_REPORTS.md), requested 14 Oct 2026: the APEX-style toolbar on every list,
the Rows selector, saved reports, Excel/CSV files, e-mailed subscriptions, and CSV import for the old-ERP migration
(four new sheets: employees, subscriptions, courses, packages).
"""
from __future__ import annotations

import io
import re
import uuid
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from app.database import SessionLocal
from app.main import app
from app.models.core import User
from app.models.ops import MigrationJob, ReportSubscription, ReportView
from app.services import ir
from app.services import migration as mig

TAG = "irt" + uuid.uuid4().hex[:6]


def _login(email: str, password: str) -> TestClient:
    c = TestClient(app)
    assert c.post("/login", data={"username": email, "password": password}, follow_redirects=False).status_code == 303
    return c


@pytest.fixture(scope="module")
def admin():
    return _login("admin@oqc.local", "Admin@12345")


@pytest.fixture(scope="module")
def teacher():
    return _login("teacher1@oqc.local", "Teacher@123")


@pytest.fixture(scope="module", autouse=True)
def _cleanup():
    yield
    s = SessionLocal()
    try:
        s.query(ReportView).filter(ReportView.name.like(f"{TAG}%")).delete(synchronize_session=False)
        s.query(ReportSubscription).filter(ReportSubscription.name.like(f"{TAG}%")).delete(synchronize_session=False)
        s.commit()
    finally:
        s.close()


def _rows_in(html: str, index: int = 0) -> int:
    return len(ir.parse_table(html, index)[1])


# ------------------------------------------------------------------------------------------------ Rows selector
def test_the_rows_choice_sets_the_page_size(admin):
    assert _rows_in(admin.get("/admin/audit").text) == 50, "the page's own size without ?rows"
    assert _rows_in(admin.get("/admin/audit?rows=100").text) == 100
    total = int(re.search(r'data-ir-pager data-total="(\d+)"', admin.get("/admin/audit").text).group(1))
    assert _rows_in(admin.get("/admin/audit?rows=all").text) == min(total, 10000)
    assert _rows_in(admin.get("/admin/audit?rows=nonsense").text) == 50


def test_the_rows_choice_never_leaks_into_the_next_request(admin):
    admin.get("/admin/audit?rows=all")
    assert _rows_in(admin.get("/admin/audit").text) == 50


# ------------------------------------------------------------------------------------------------ the toolbar on pages
def test_staff_get_the_toolbar_and_the_import_button_where_they_may_import(admin, teacher):
    page = admin.get("/clients").text
    assert "table" in page and "data-ir" in page and "/static/js/report.js" in page
    assert 'data-ir-import="clients"' in page, "the admin may import families"
    t = teacher.get("/students").text
    assert "/static/js/report.js" in t and "data-ir-import" not in t, "no import without migration permission"
    parent = _login("parent1@oqc.local", "Parent@123")
    assert "/static/js/report.js" not in parent.get("/portal").text, "families do not get staff tools"


def test_every_list_page_has_its_import_sheet():
    assert mig.entity_for_path("/clients") == "clients"
    assert mig.entity_for_path("/hr/employees/") == "employees"
    assert mig.entity_for_path("/finance/subscriptions") == "subscriptions"
    assert mig.entity_for_path("/academics/config/packages") == "packages"
    assert mig.entity_for_path("/home") is None
    assert set(mig.PAGE_ENTITIES.values()) <= set(mig.ENTITIES)


# ------------------------------------------------------------------------------------------------ the server's copy of the browser rules
def test_reading_a_page_and_applying_saved_settings(admin):
    headers, rows = ir.parse_table(admin.get("/clients?rows=all").text)
    assert headers[:3] == ["ID", "Client Name", "Status"] and rows
    status = headers.index("Status")
    config = {"filters": [{"col": status, "op": "eq", "value": "Regular"}], "sort": [{"col": 1, "dir": "asc"}],
              "hidden": [headers.index("Email")], "actions": [len(headers) - 1]}
    h2, r2 = ir.apply_config(headers, rows, config)
    assert "Email" not in h2 and "" not in h2, "hidden and button columns are left out"
    assert r2 and all(r[h2.index("Status")] == "Regular" for r in r2)
    names = [r[h2.index("Client Name")] for r in r2]
    assert names == sorted(names, key=str.lower)


def test_number_and_rule_semantics():
    assert ir.number("PKR 1,250.50") == 1250.5 and ir.number("£12") == 12 and ir.number("45%") == 45
    assert ir.number("C-00045") is None and ir.number("") is None
    assert ir.matches("1,200", "gt", "1000") and not ir.matches("900", "gt", "1000")
    assert ir.matches("Trial", "eq", "trial") and ir.matches("", "empty", "") and ir.matches("abc", "starts", "AB")


# ------------------------------------------------------------------------------------------------ saved reports
def test_saved_reports_are_private_unless_shared(admin):
    r = admin.post("/ir/views", json={"path": "/clients", "name": f"{TAG} mine", "config": {"filters": []}})
    assert r.status_code == 200 and r.json()["mine"]
    shared = admin.post("/ir/views", json={"path": "/clients", "name": f"{TAG} shared", "config": {}, "shared": True}).json()
    hr = _login("hr@oqc.local", "People@123")
    names = {v["name"] for v in hr.get("/ir/views?path=/clients").json()["views"]}
    assert f"{TAG} shared" in names and f"{TAG} mine" not in names
    assert hr.post(f"/ir/views/{shared['id']}/delete").status_code == 404, "only the owner deletes"
    assert admin.post(f"/ir/views/{shared['id']}/delete").status_code == 200


def test_a_report_cannot_be_saved_for_a_page_the_user_cannot_open(teacher):
    assert teacher.post("/ir/views", json={"path": "/hr/grievances", "name": f"{TAG} x", "config": {}}).status_code == 403
    assert teacher.post("/ir/views", json={"path": "//evil.example", "name": f"{TAG} x"}).status_code == 400


def test_families_cannot_use_the_report_endpoints():
    parent = _login("parent1@oqc.local", "Parent@123")
    assert parent.get("/ir/views?path=/portal").status_code == 403
    assert parent.post("/ir/export", json={"headers": ["a"], "rows": [["1"]]}).status_code == 403


# ------------------------------------------------------------------------------------------------ files
def test_excel_and_csv_files(admin):
    r = admin.post("/ir/export", json={"name": "Client List", "headers": ["Code", "Fee"], "rows": [["007", "PKR 1,250.50"], ["8", "n/a"]]})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["Code", "Fee"]
    assert ws["A2"].value == "007", "codes with a leading zero stay text"
    assert ws["B2"].value == 1250.5, "amounts become numbers"
    c = admin.post("/ir/export", json={"name": "x", "headers": ["A"], "rows": [["ع"]], "format": "csv"})
    assert c.content.startswith("﻿".encode()) and "ع" in c.content.decode("utf-8-sig")


# ------------------------------------------------------------------------------------------------ subscriptions
def test_a_subscription_mails_what_the_subscriber_sees(admin, monkeypatch):
    sent = []
    monkeypatch.setattr("app.services.integrations.send_email", lambda db, to, subject, body, html=False, attachments=None:
                        sent.append((to, subject, attachments)) or {"simulated": True})
    headers, _ = ir.parse_table(admin.get("/clients").text)
    bad = admin.post("/ir/subscriptions", json={"path": "/clients", "name": f"{TAG} sub", "emails": "nope", "frequency": "weekly"})
    assert bad.status_code == 400
    sub = admin.post("/ir/subscriptions", json={
        "path": "/clients", "query": "?page=3&rows=50", "name": f"{TAG} sub", "emails": "ops@example.com; boss@example.com",
        "frequency": "weekly", "config": {"filters": [{"col": headers.index("Status"), "op": "eq", "value": "Trial"}],
                                          "actions": [len(headers) - 1]}}).json()
    assert sub["next_run_at"] and sub["emails"] == "ops@example.com, boss@example.com"
    done = admin.post(f"/ir/subscriptions/{sub['id']}/send").json()
    assert done["last_status"].startswith("Sent"), done["last_status"]
    assert len(sent) == 2 and sent[0][2][0][0].endswith(".csv")
    body = sent[0][2][0][1].decode("utf-8-sig").splitlines()
    assert body[0].split(",")[:3] == ["ID", "Client Name", "Status"]
    assert len(body) > 1 and all(",Trial," in line for line in body[1:])
    s = SessionLocal()
    try:
        row = s.get(ReportSubscription, sub["id"])
        assert row.query is None or "page=" not in row.query, "the page and rows of the moment are not kept"
        row.next_run_at = datetime.utcnow() - timedelta(minutes=1)
        s.commit()
        result = ir.run_due(s, now=datetime.utcnow())
        s.commit()
        assert result["sent"] >= 1
        s.refresh(row)
        assert row.next_run_at > datetime.utcnow() + timedelta(days=6)
    finally:
        s.close()


def test_a_subscription_fails_cleanly_when_access_is_gone(monkeypatch):
    monkeypatch.setattr("app.services.integrations.send_email", lambda *a, **k: {"simulated": True})
    s = SessionLocal()
    try:
        teacher = s.query(User).filter(User.email == "teacher1@oqc.local").one()
        sub = ReportSubscription(user_id=teacher.id, path="/clients", name=f"{TAG} lost", emails="t@example.com",
                                 frequency="daily", config={}, next_run_at=datetime.utcnow() - timedelta(minutes=1))
        s.add(sub)
        s.commit()
        ir.run_due(s)
        s.commit()
        s.refresh(sub)
        assert sub.last_status.startswith("Failed") and "403" in sub.last_status
    finally:
        s.close()


# ------------------------------------------------------------------------------------------------ CSV import: the four new sheets
def _import(entity: str, text: str, user_id: int) -> tuple[MigrationJob, dict]:
    s = SessionLocal()
    path = mig.save_upload(f"{TAG}-{entity}.csv", text.encode("utf-8"))
    headers, rows = mig.parse_file(path)
    job = MigrationJob(name=f"{TAG} {entity}", entity=entity, source_system="test", file_path=str(path), status="uploaded",
                       records_total=len(rows), field_mapping=mig.suggest_mapping(entity, headers), created_by_id=user_id, errors=[])
    s.add(job)
    s.flush()
    check = mig.validate_job(s, job)
    result = mig.import_job(s, job, s.get(User, user_id))
    s.commit()
    result["validation"] = check
    s.close()
    return job, result


def test_courses_packages_employees_and_subscriptions_import():
    from app.models.academic import Course, Package
    from app.models.finance import Subscription
    from app.models.people import Employee, Student
    s = SessionLocal()
    admin = s.query(User).filter(User.email == "admin@oqc.local").one()
    student = s.query(Student).filter(Student.client_id.isnot(None)).order_by(Student.id).first()
    code = TAG.upper()[:8]
    s.close()
    try:
        _, r = _import("courses", f"code,name,fee\n{code},Test course {TAG},4000\n{code},Duplicate {TAG},1\n", admin.id)
        assert r["imported"] == 1 and r["duplicates"] == 1, r
        _, r = _import("packages", f"name,course,sessions_per_week,price,currency\n{TAG} pack,{code},3,45,GBP\n", admin.id)
        assert r["imported"] == 1, r
        _, r = _import("employees", "employee_code,full_name,designation,department,email,join_date,base_salary\n"
                                    f"{code}-1,Imported Person,Teacher,academics,{TAG}@example.com,03/01/2025,55000\n"
                                    f"{code}-2,Wrong Dept,Clerk,no-such-department,,,\n", admin.id)
        assert r["imported"] == 1 and r["skipped"] == 1, r
        _, r = _import("subscriptions", "subscription_code,student_code,package,price,currency,days,start_date\n"
                                        f"SUB-{code},{student.student_code},{TAG} pack,45,GBP,\"Mon, Wed, Fri\",2025-09-01\n"
                                        f"SUB-{code}X,NO-SUCH-STUDENT,,10,GBP,,\n", admin.id)
        assert r["imported"] == 1 and r["skipped"] == 1, r
        s = SessionLocal()
        sub = s.query(Subscription).filter(Subscription.subscription_code == f"SUB-{code}").one()
        assert sub.days_of_week == [0, 2, 4] and sub.package.name == f"{TAG} pack" and sub.course.code == code
        emp = s.query(Employee).filter(Employee.employee_code == f"{code}-1").one()
        assert emp.department.code == "academics" and float(emp.base_salary) == 55000 and emp.join_date.year == 2025
        s.close()
        # importing the same file again reports duplicates instead of creating them twice
        _, again = _import("courses", f"code,name\n{code},Test course {TAG}\n", admin.id)
        assert again["imported"] == 0 and again["duplicates"] == 1
    finally:
        s = SessionLocal()
        s.query(Subscription).filter(Subscription.subscription_code.like(f"SUB-{code}%")).delete(synchronize_session=False)
        s.query(Employee).filter(Employee.employee_code.like(f"{code}-%")).delete(synchronize_session=False)
        s.query(Package).filter(Package.name == f"{TAG} pack").delete(synchronize_session=False)
        s.query(Course).filter(Course.code == code).delete(synchronize_session=False)
        s.query(MigrationJob).filter(MigrationJob.name.like(f"{TAG}%")).delete(synchronize_session=False)
        s.commit()
        s.close()


def test_the_import_page_opens_on_the_sheet_of_the_list(admin):
    page = admin.get("/admin/migration?entity=employees&back=/hr/employees").text
    assert re.search(r'<option value="employees"\s+selected', page)
    assert 'name="back" value="/hr/employees"' in page
    unsafe = admin.get("/admin/migration?entity=employees&back=//evil.example").text
    assert 'value="//evil.example"' not in unsafe


# ------------------------------------------------------------------------------------------------ old-ERP files (16 Oct 2026)
def test_old_erp_headings_map_to_the_right_fields():
    """The college's first imports skipped rows: headings such as "Client Name" or "Net Amount" were not mapped."""
    m = mig.suggest_mapping("clients", ["ID", "Client Name", "Status", "Registration Date", "Email", "Cell", "WhatsApp", "Status Remarks"])
    assert m == {"client_code": "ID", "full_name": "Client Name", "status": "Status", "joined_at": "Registration Date",
                 "email": "Email", "phone": "Cell", "whatsapp": "WhatsApp", "notes": "Status Remarks"}
    inv = mig.suggest_mapping("invoices", ["ID", "Client", "Invoice Date", "Total Amount", "Discount Amount", "Net Amount", "Status"])
    assert inv["total"] == "Net Amount" and inv["client_name"] == "Client" and inv["issue_date"] == "Invoice Date"
    pay = mig.suggest_mapping("payments", ["ID", "Receipt Date", "Client", "Total Amount", "Ref No", "Payment Mode"])
    assert pay["received_at"] == "Receipt Date" and pay["amount"] == "Total Amount" and pay["method"] == "Payment Mode"
    pk = mig.suggest_mapping("packages", ["ID", "Description", "Maximum Days", "Session", "Price", "Course"])
    assert pk["name"] == "Description", "a required field outranks an optional one for the same heading"
    emp = mig.suggest_mapping("employees", ["ID", "Name", "Designation", "Cell", "Basic Salary", "Joining Date"])
    assert emp["full_name"] == "Name" and emp["base_salary"] == "Basic Salary" and emp["join_date"] == "Joining Date"
    # a file that already uses our field names is unchanged
    assert mig.suggest_mapping("clients", ["client_code", "full_name", "email"]) == {"client_code": "client_code", "full_name": "full_name", "email": "email"}
    assert mig.coerce("30 Minutes", "int") == (30, None)


def test_a_bad_optional_value_is_left_out_and_reported():
    from app.models.people import Client
    s = SessionLocal()
    admin = s.query(User).filter(User.email == "admin@oqc.local").one()
    s.close()
    code = f"OLD-{TAG[-5:].upper()}"
    try:
        job, r = _import("clients", f"ID,Client Name,Email,Cell\n{code},Bad Row {TAG},bad-email,123\n", admin.id)
        assert r["imported"] == 1, r
        s = SessionLocal()
        c = s.query(Client).filter(Client.client_code == code).one()
        assert c.email is None and c.phone is None, "never stored as it was"
        notes = [e["message"] for e in s.get(MigrationJob, job.id).errors]
        assert any("Imported without email" in n for n in notes) and any("Imported without phone" in n for n in notes)
        s.close()
    finally:
        s = SessionLocal()
        s.query(Client).filter(Client.client_code == code).delete(synchronize_session=False)
        s.query(MigrationJob).filter(MigrationJob.name.like(f"{TAG}%")).delete(synchronize_session=False)
        s.commit()
        s.close()


def test_invoices_find_the_family_by_name_when_the_file_has_no_code():
    from app.models.finance import Invoice, LedgerEntry
    from app.models.people import Client
    s = SessionLocal()
    admin = s.query(User).filter(User.email == "admin@oqc.local").one()
    fam = Client(client_code=f"C-{TAG[-5:].upper()}", full_name=f"Unique Family {TAG}", country="UK", currency="GBP", status="active")
    s.add(fam)
    s.commit()
    fam_id = fam.id
    s.close()
    try:
        _, r = _import("invoices", f"ID,Client,Invoice Date,Net Amount\nINV-{TAG},Unique Family {TAG},01/09/2026,45\n", admin.id)
        assert r["imported"] == 1, r
        s = SessionLocal()
        assert s.query(Invoice).filter(Invoice.invoice_number == f"INV-{TAG}", Invoice.client_id == fam_id).count() == 1
        s.close()
    finally:
        s = SessionLocal()
        inv_ids = [i for (i,) in s.query(Invoice.id).filter(Invoice.invoice_number == f"INV-{TAG}")]
        s.query(LedgerEntry).filter(LedgerEntry.client_id == fam_id).delete(synchronize_session=False)
        s.query(Invoice).filter(Invoice.id.in_(inv_ids or [-1])).delete(synchronize_session=False)
        s.query(Client).filter(Client.id == fam_id).delete(synchronize_session=False)
        s.query(MigrationJob).filter(MigrationJob.name.like(f"{TAG}%")).delete(synchronize_session=False)
        s.commit()
        s.close()


def test_every_sheet_card_has_an_import_button(admin):
    page = admin.get("/admin/migration").text
    for key in mig.ENTITIES:
        assert f'href="/admin/migration?entity={key}"' in page, key
    assert "Courses → Packages → Employees" in page
    assert mig.ENTITIES["leads"]["icon"] == "filter", "funnel is not in the bundled icon set"


def test_teachers_import_with_their_courses_and_verification():
    """17 Oct 2026: teachers from the old ERP already teach; the import can bring their courses and mark them verified."""
    from app.models.core import AuditEvent
    from app.models.people import Teacher
    s = SessionLocal()
    admin = s.query(User).filter(User.email == "admin@oqc.local").one()
    teacher_user = s.query(User).filter(User.email == "teacher1@oqc.local").one()
    s.close()
    code = f"TV{TAG[-5:].upper()}"
    try:
        _, r = _import("teachers", f"Code,Teacher,Courses,Verified\n{code}1,Imported Verified {TAG},\"NAZRA, hifz\",yes\n{code}2,Imported Pending {TAG},NAZRA,no\n", admin.id)
        assert r["imported"] == 2, r
        s = SessionLocal()
        yes = s.query(Teacher).filter(Teacher.teacher_code == f"{code}1").one()
        no = s.query(Teacher).filter(Teacher.teacher_code == f"{code}2").one()
        assert yes.is_verified is True and yes.courses == ["NAZRA", "HIFZ"]
        assert no.is_verified is False and no.courses == ["NAZRA"]
        assert s.query(AuditEvent).filter(AuditEvent.entity_type == "Teacher", AuditEvent.entity_id == yes.id, AuditEvent.action == "approve").count() == 1
        s.close()
        # someone who may not approve teachers cannot import them as verified
        _, r2 = _import("teachers", f"Code,Teacher,Verified\n{code}3,No Right {TAG},yes\n", teacher_user.id)
        assert r2["imported"] == 0 and r2["skipped"] == 1
    finally:
        s = SessionLocal()
        ids = [i for (i,) in s.query(Teacher.id).filter(Teacher.teacher_code.like(f"{code}%"))]
        s.query(AuditEvent).filter(AuditEvent.entity_type == "Teacher", AuditEvent.entity_id.in_(ids or [-1])).delete(synchronize_session=False)
        s.query(Teacher).filter(Teacher.id.in_(ids or [-1])).delete(synchronize_session=False)
        s.query(MigrationJob).filter(MigrationJob.name.like(f"{TAG}%")).delete(synchronize_session=False)
        s.commit()
        s.close()
