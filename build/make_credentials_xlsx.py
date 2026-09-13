"""Generate the account & access workbook for the live deployment.

    .venv/Scripts/python.exe build/make_credentials_xlsx.py

Reads live_accounts.json (pulled from the deployed API) and writes
"Online Quran College - User Accounts.xlsx".
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
SITE = "https://oqc.onrender.com"
ADMIN_PW = "UEs6t8gS4+3/7GfWCNW0Wx73585GnvyU0qLjsIZvVgA="
DEMO_PW = "jz97UqlDbrIStPPI91uFthratTrtqM5yMcAG/VqO5XY="

FONT = "Arial"
NAVY = "0F172A"
BRAND = "0F766E"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
BAND = PatternFill("solid", fgColor="F1F5F9")
WARN = PatternFill("solid", fgColor="FEF3C7")
THIN = Side(style="thin", color="CBD5E1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# username, name, role, portal, group, access summary
STAFF = [
    ("admin@oqc.local", "Muhammad Yousuf (CEO)", "Super Admin / CEO", "Admin", "Leadership",
     "Everything. CEO command centre, anti-poaching flags, staff eNPS, every approval tier"),
    ("sysadmin@oqc.local", "Bilal Ahmed", "System Administrator", "Admin", "Leadership",
     "Users, roles, settings, integrations, API keys, backups, migration, audit log"),
    ("hr@oqc.local", "Sana Malik", "HOD People & Culture", "Admin", "Leadership",
     "Employees, payroll, recruitment, teacher grading, confidential grievances"),
    ("finance@oqc.local", "Imran Qureshi", "HOD Finance", "Admin", "Leadership",
     "Invoices, payments, accounts, P&L, budgets, discount and scholarship approval"),
    ("academics@oqc.local", "Qari Abdul Rehman", "HOD Academics", "Admin", "Leadership",
     "Curriculum, lesson plans, evaluations, monthly tests, certificates, scheduling"),
    ("qa@oqc.local", "Hafiz Usman Tariq", "HOD Quality Assurance", "Admin", "Leadership",
     "QA queue, scorecards, corrective actions, AI class monitoring, safeguarding view"),
    ("tech@oqc.local", "Faisal Khan", "HOD Technology", "Admin", "Leadership",
     "Infrastructure, integrations, security centre, API health, audit log"),
    ("marketing@oqc.local", "Ayesha Siddiqui", "HOD Marketing", "Admin", "Leadership",
     "Leads, campaigns, attribution, CPL and CAC, WhatsApp inbox, referrals"),
    ("manager@oqc.local", "Kamran Ali", "Manager (Morning Group)", "Admin", "Operations",
     "Students, scheduling, retention, cases, KPIs. Approves discounts up to 20%"),
    ("manager2@oqc.local", "Nadia Hussain", "Manager (Night Group)", "Admin", "Operations",
     "Same as morning manager, scoped to the night shift"),
    ("supervisor@oqc.local", "Zubair Shah", "Supervisor (Evening)", "Admin", "Operations",
     "Live class board for their own teachers only. Marks missed, absent, cancelled"),
    ("supervisor2@oqc.local", "Rabia Noor", "Supervisor (Night)", "Admin", "Operations",
     "Same as evening supervisor, scoped to their own assigned teachers"),
    ("coordinator@oqc.local", "Ustadha Fatima Zahra", "Academic Coordinator", "Admin", "Operations",
     "Curriculum, schedules, lesson plans, evaluations, certificates"),
    ("leadgen@oqc.local", "Saad Farooq", "Lead Generator", "Admin", "Commercial",
     "Creates leads and works the WhatsApp inbox. Sees only leads they generated"),
    ("closer@oqc.local", "Omar Javed", "Lead Closer", "Admin", "Commercial",
     "Full pipeline, trials, registration, converts leads into enrolled students"),
    ("billing@oqc.local", "Hina Baig", "Billing Representative", "Admin", "Commercial",
     "Invoices, payments, receipts, client ledger. Cannot approve discounts"),
    ("accountant@oqc.local", "Tariq Mehmood", "Accountant", "Admin", "Commercial",
     "Chart of accounts, journals, expenses, P&L, cash flow, currencies"),
    ("qaofficer@oqc.local", "Maryam Iqbal", "QA Officer", "Admin", "Specialist",
     "Carries out class reviews and scorecards, raises corrective actions"),
    ("hrofficer@oqc.local", "Adeel Raza", "HR Officer", "Admin", "Specialist",
     "Attendance, leave, recruitment, violations. Limited grievance visibility"),
    ("auditor@oqc.local", "External Auditor", "External Auditor", "Admin", "Specialist",
     "Read-only across the whole platform. Cannot change anything (verified)"),
]


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 26


def write_sheet(ws, headers, rows, widths, start=1, banded=True):
    for i, h in enumerate(headers, start=1):
        ws.cell(row=start, column=i, value=h)
    style_header(ws, start, len(headers))
    for r, data in enumerate(rows, start=start + 1):
        for c, val in enumerate(data, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(c == len(headers)))
            cell.border = BOX
            if banded and (r - start) % 2 == 0:
                cell.fill = BAND
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def main() -> None:
    live = json.loads((ROOT / "live_accounts.json").read_text(encoding="utf8"))
    wb = Workbook()

    # ---------------------------------------------------------------- Read me
    ws = wb.active
    ws.title = "Read me"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Online Quran College - Digital Operating System"
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    ws["A2"] = "User accounts and access levels"
    ws["A2"].font = Font(name=FONT, size=11, color=BRAND)

    info = [
        ("Live site", SITE),
        ("Super admin username", "admin@oqc.local"),
        ("Super admin password", ADMIN_PW),
        ("Password for all other accounts", DEMO_PW),
        ("", ""),
        ("Total accounts", "122"),
        ("Staff (admin portal)", "20"),
        ("Teachers (teacher portal)", str(len(live["teachers"]))),
        ("Parents / clients (client portal)", str(len(live["clients"]))),
        ("Students (student portal)", str(len(live["students"]))),
    ]
    r = 4
    for label, value in info:
        if not label:
            r += 1
            continue
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, size=10, bold=True)
        v = ws.cell(row=r, column=2, value=value)
        v.font = Font(name=FONT, size=10)
        if "password" in label.lower():
            v.font = Font(name="Consolas", size=10, bold=True)
            v.fill = WARN
        r += 1

    notes = [
        "",
        "Notes",
        "1. The super admin account is prompted to set a new password on first sign-in. The password above stops working after that.",
        "2. Every other account shares one password so a client walkthrough is easy. Change this before real use.",
        "3. Access is enforced by the system, not by convention. A supervisor sees only their own teachers, and a",
        "    parent cannot open another family's child. This was tested, not assumed.",
        "4. Five failed sign-in attempts lock an account for 15 minutes. If that happens, wait rather than retrying.",
        "5. The free Render instance sleeps after about 15 minutes idle, so the first visit can take 30 to 50 seconds.",
        "6. Passwords live in Render under service oqc, Environment tab (ADMIN_PASSWORD and DEMO_PASSWORD).",
    ]
    for line in notes:
        c = ws.cell(row=r, column=1, value=line)
        c.font = Font(name=FONT, size=10, bold=(line == "Notes"), color=NAVY if line == "Notes" else "334155")
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 62

    # ---------------------------------------------------------------- Staff
    ws = wb.create_sheet("Staff accounts")
    rows = [[u, n, role, portal, group, DEMO_PW if u != "admin@oqc.local" else ADMIN_PW, access]
            for u, n, role, portal, group, access in STAFF]
    write_sheet(ws, ["Username", "Name", "Role", "Portal", "Group", "Password", "What they can access"],
                rows, [24, 24, 26, 9, 13, 46, 74])
    for row in range(2, len(rows) + 2):
        ws.cell(row=row, column=6).font = Font(name="Consolas", size=9)
    ws.cell(row=2, column=6).fill = WARN  # highlight the super admin password

    # ---------------------------------------------------------------- Teachers
    ws = wb.create_sheet("Teachers")
    rows = []
    for i, t in enumerate(sorted(live["teachers"], key=lambda x: x["teacher_code"]), start=1):
        rows.append([f"teacher{i}@oqc.local", t["full_name"], t["teacher_code"], t["shift"].title(),
                     f"Grade {t['grade']}", ", ".join(t.get("courses") or []),
                     "Yes" if t.get("is_verified") else "No - cannot take live classes", DEMO_PW,
                     "Teacher portal: own schedule, classes, lesson plans, students, QA feedback, payslips"])
    write_sheet(ws, ["Username", "Name", "Code", "Shift", "Grade", "Courses", "Background check",
                     "Password", "What they can access"],
                rows, [22, 26, 11, 11, 10, 26, 26, 46, 60])
    for row in range(2, len(rows) + 2):
        ws.cell(row=row, column=8).font = Font(name="Consolas", size=9)

    # ---------------------------------------------------------------- Parents
    ws = wb.create_sheet("Parents (clients)")
    rows = []
    for i, c in enumerate(sorted(live["clients"], key=lambda x: x["client_code"]), start=1):
        rows.append([f"parent{i}@oqc.local", c["full_name"], c["client_code"], c.get("country", ""),
                     c.get("city", ""), c.get("currency", ""), DEMO_PW,
                     "Client portal: own children only. Schedule, progress, result cards, invoices, referrals"])
    write_sheet(ws, ["Username", "Name", "Code", "Country", "City", "Currency", "Password",
                     "What they can access"],
                rows, [22, 26, 11, 18, 15, 10, 46, 66])
    for row in range(2, len(rows) + 2):
        ws.cell(row=row, column=7).font = Font(name="Consolas", size=9)

    # ---------------------------------------------------------------- Students
    ws = wb.create_sheet("Students")
    rows = []
    for i, s in enumerate(sorted(live["students"], key=lambda x: x["student_code"]), start=1):
        rows.append([f"student{i}@oqc.local", s["full_name"], s["student_code"], s.get("status", ""),
                     s.get("age", ""), s.get("level") or "", DEMO_PW,
                     "Student portal: own classes, progress, Arabic lesson view, results, certificates"])
    write_sheet(ws, ["Username", "Name", "Code", "Status", "Age", "Level", "Password",
                     "What they can access"],
                rows, [22, 26, 11, 12, 7, 24, 46, 64])
    for row in range(2, len(rows) + 2):
        ws.cell(row=row, column=7).font = Font(name="Consolas", size=9)

    # ---------------------------------------------------------------- Demo route
    ws = wb.create_sheet("Demo route")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Suggested order for a client walkthrough"
    ws["A1"].font = Font(name=FONT, size=13, bold=True, color=NAVY)
    steps = [
        ("1", "admin@oqc.local", "CEO Command Center", "Institutional health score, revenue, churn, AI insights"),
        ("2", "admin@oqc.local", "Academic Home", "Live counters for today: pending, done, missed, absent"),
        ("3", "supervisor@oqc.local", "Supervisor Live", "Real-time class board. Note it shows only their own teachers"),
        ("4", "teacher1@oqc.local", "Teacher portal", "Today's classes, start a class, lesson plan, monthly test scoring"),
        ("5", "parent1@oqc.local", "Parent portal", "The same class seen by the family: progress, result card, invoice"),
        ("6", "admin@oqc.local", "Finance", "Discount ladder: above 35% is refused outright"),
        ("7", "admin@oqc.local", "QA and AI monitoring", "Class analysis with confidence and human review required"),
        ("8", "auditor@oqc.local", "Any page", "Read-only proof: every mutating action is refused"),
    ]
    write_sheet(ws, ["Step", "Sign in as", "Where to go", "What it shows"],
                [list(s) for s in steps], [7, 24, 26, 74], start=3)
    wb.save(ROOT / "Online Quran College - User Accounts.xlsx")
    print("Written: Online Quran College - User Accounts.xlsx")


if __name__ == "__main__":
    main()
