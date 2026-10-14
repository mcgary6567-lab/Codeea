"""Interactive reports (docs/INTERACTIVE_REPORTS.md): saved views, Excel export and e-mailed subscriptions.

The toolbar itself runs in the browser (app/static/js/report.js) on every list table. The server keeps what has to
outlive the page: saved views, subscriptions, and the files. A subscription renders its page *as the subscriber*
(an internal request with a short-lived session), so the file carries exactly what that person may see, then
applies the same column, filter and sort settings the browser applies.
"""
from __future__ import annotations

import csv
import io
import re
import secrets
from datetime import datetime, timedelta
from html import unescape
from typing import Optional

from sqlalchemy.orm import Session

from app.models.core import User, UserSession
from app.models.ops import ReportSubscription

FREQUENCIES = {"daily": timedelta(days=1), "weekly": timedelta(days=7), "monthly": timedelta(days=30)}
OPERATORS = ("contains", "not_contains", "eq", "neq", "gt", "lt", "gte", "lte", "starts", "empty", "not_empty")
_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


# ----------------------------------------------------------------------------- reading a rendered table
def _text(fragment: str) -> str:
    fragment = re.sub(r"<(script|style)\b.*?</\1>", " ", fragment, flags=re.S | re.I)
    fragment = re.sub(r"<[^>]+>", " ", fragment)
    return re.sub(r"\s+", " ", unescape(fragment)).strip()


def parse_table(html: str, index: int = 0) -> tuple[list[str], list[list[str]]]:
    """Headers and data rows of the ``index``-th interactive table (``data-ir``) of a rendered page."""
    tables = re.findall(r"<table\b[^>]*\bdata-ir\b[^>]*>(.*?)</table>", html, flags=re.S | re.I)
    if index >= len(tables):
        return [], []
    body = tables[index]
    head = re.search(r"<thead\b.*?>(.*?)</thead>", body, flags=re.S | re.I)
    headers = [_text(h) for h in re.findall(r"<th\b[^>]*>(.*?)</th>", head.group(1) if head else "", flags=re.S | re.I)]
    tbody = re.search(r"<tbody\b[^>]*>(.*)</tbody>", body, flags=re.S | re.I)
    rows: list[list[str]] = []
    for tr in re.findall(r"<tr\b[^>]*>(.*?)</tr>", tbody.group(1) if tbody else "", flags=re.S | re.I):
        cells = re.findall(r"<td\b([^>]*)>(.*?)</td>", tr, flags=re.S | re.I)
        if len(cells) == 1 and "colspan" in cells[0][0]:
            continue  # the "nothing here" row
        rows.append([_text(c[1]) for c in cells])
    return headers, rows


def action_columns(headers: list[str], config: dict | None = None) -> set[int]:
    """Columns that hold buttons, not data. The browser decides (it can see the controls) and saves them as
    ``actions``; without that, only columns with no heading count."""
    if config and isinstance(config.get("actions"), list):
        return {int(i) for i in config["actions"]}
    return {i for i, h in enumerate(headers) if not h.strip()}


# ----------------------------------------------------------------------------- the browser's settings, in Python
def number(value: str) -> Optional[float]:
    """'PKR 1,250.50', '£12', '45%' -> float; None when the cell is not a number. Mirrors report.js num()."""
    s = (value or "").strip()
    if not s or not re.search(r"\d", s):
        return None
    # only a currency in front ("PKR 1,250", "£12") or a unit behind ("45%", "12 GBP"); never the letters of a
    # code: "C-00045" is not -45
    stripped = re.sub(r"^(?:[A-Za-z]{3}\s+|[£$€₨]\s*)|\s*(?:%|[A-Za-z]{3})$", "", s).replace(",", "")
    return float(stripped) if re.fullmatch(r"-?\d+(\.\d+)?", stripped) else None


def matches(cell: str, op: str, value: str) -> bool:
    a, b = (cell or "").strip(), (value or "").strip()
    al, bl = a.lower(), b.lower()
    if op == "empty":
        return a == ""
    if op == "not_empty":
        return a != ""
    if op == "contains":
        return bl in al
    if op == "not_contains":
        return bl not in al
    if op == "starts":
        return al.startswith(bl)
    if op in ("eq", "neq"):
        na, nb = number(a), number(b)
        same = (na == nb) if na is not None and nb is not None else al == bl
        return same if op == "eq" else not same
    na, nb = number(a), number(b)
    x, y = (na, nb) if na is not None and nb is not None else (al, bl)
    return {"gt": x > y, "lt": x < y, "gte": x >= y, "lte": x <= y}.get(op, True)


def apply_config(headers: list[str], rows: list[list[str]], config: dict) -> tuple[list[str], list[list[str]]]:
    """Filters, search, sort and visible columns as saved by the toolbar; action columns are always dropped."""
    config = config or {}
    out = rows
    search = (config.get("search") or "").strip().lower()
    if search:
        out = [r for r in out if any(search in c.lower() for c in r)]
    for f in config.get("filters") or []:
        col, op, val = int(f.get("col", -1)), f.get("op", "contains"), str(f.get("value", ""))
        if 0 <= col < len(headers):
            out = [r for r in out if matches(r[col] if col < len(r) else "", op, val)]
    for s in reversed(config.get("sort") or []):
        col, desc = int(s.get("col", -1)), s.get("dir") == "desc"
        if 0 <= col < len(headers):
            def key(r, c=col):
                v = r[c] if c < len(r) else ""
                n = number(v)
                return (0, n, "") if n is not None else (1, 0, v.lower())
            out = sorted(out, key=key, reverse=desc)
    hidden = {int(i) for i in config.get("hidden") or []} | action_columns(headers, config)
    keep = [i for i in range(len(headers)) if i not in hidden]
    return [headers[i] for i in keep], [[r[i] if i < len(r) else "" for i in keep] for r in out]


# ----------------------------------------------------------------------------- files
def csv_bytes(headers: list[str], rows: list[list[str]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # the BOM makes Excel read UTF-8 (Arabic, Urdu, £)


def xlsx_bytes(title: str, headers: list[str], rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\[\]:*?/\\]", " ", title or "Report").strip() or "Report")[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8F1FB")
    for r in rows:
        ws.append([(number(v) if number(v) is not None and not re.match(r"^0\d", v.strip()) else v) for v in r])
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if i - 1 < len(r)])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(60, max(8, width + 2))
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def slug(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", name or "report").strip("-").lower()[:60] or "report"


# ----------------------------------------------------------------------------- subscriptions
def clean_emails(raw: str) -> list[str]:
    emails = [e.strip().lower() for e in re.split(r"[,;\s]+", raw or "") if e.strip()]
    bad = [e for e in emails if not _EMAIL.match(e)]
    if bad:
        raise ValueError(f"Not an e-mail address: {', '.join(bad[:3])}")
    if not emails:
        raise ValueError("Give at least one e-mail address.")
    if len(emails) > 10:
        raise ValueError("At most 10 recipients per subscription.")
    return emails


def next_run(frequency: str, after: Optional[datetime] = None) -> datetime:
    base = (after or datetime.utcnow()).replace(minute=0, second=0, microsecond=0)
    return base + FREQUENCIES.get(frequency, FREQUENCIES["weekly"])


def render_as(user: User, path: str, query: str = "") -> tuple[int, str]:
    """GET a page inside the app as ``user``: a session that lives for this one request, revoked straight after."""
    from fastapi.testclient import TestClient

    from app.core.security import SESSION_COOKIE, create_access_token
    from app.database import SessionLocal
    from app.main import app
    db = SessionLocal()
    jti = secrets.token_hex(16)
    try:
        db.add(UserSession(user_id=user.id, token_jti=jti, ip="report-subscription", user_agent="report subscription",
                           expires_at=datetime.utcnow() + timedelta(minutes=10)))
        db.commit()
        client = TestClient(app, base_url="https://internal")
        client.cookies.set(SESSION_COOKIE, create_access_token(user.id, jti=jti, minutes=10))
        sep = "&" if query else ""
        r = client.get(f"{path}?{query}{sep}rows=all", follow_redirects=False)
        return r.status_code, r.text
    finally:
        s = db.query(UserSession).filter(UserSession.token_jti == jti).first()
        if s is not None:
            s.revoked = True
            db.commit()
        db.close()


def build(user: User, sub: ReportSubscription) -> tuple[list[str], list[list[str]]]:
    status, html = render_as(user, sub.path, sub.query or "")
    if status != 200:
        raise ValueError(f"the page answered {status} (permission changed, or a password change is due)")
    headers, rows = parse_table(html, int((sub.config or {}).get("table", 0)))
    if not headers:
        raise ValueError("the page no longer has that table")
    return apply_config(headers, rows, sub.config or {})


def send(db: Session, sub: ReportSubscription) -> dict:
    from app.services.integrations import send_email
    user = db.get(User, sub.user_id)
    if user is None or not user.is_active:
        raise ValueError("the subscriber's account is inactive")
    headers, rows = build(user, sub)
    stamp = datetime.utcnow().strftime("%Y-%m-%d")
    filename = f"{slug(sub.name)}-{stamp}.csv"
    body = (f"{sub.name}\n{len(rows)} row(s), {sub.frequency} subscription of {user.full_name}.\n"
            f"Page: {sub.path}{'?' + sub.query if sub.query else ''}\n\nThe CSV is attached.")
    for to in clean_emails(sub.emails):
        send_email(db, to, f"[OQC report] {sub.name} - {stamp}", body, attachments=[(filename, csv_bytes(headers, rows), "text/csv")])
    return {"rows": len(rows), "file": filename}


def run_due(db: Session, now: Optional[datetime] = None) -> dict:
    now = now or datetime.utcnow()
    due = (db.query(ReportSubscription).filter(ReportSubscription.is_active.is_(True), ReportSubscription.next_run_at <= now)
           .order_by(ReportSubscription.next_run_at).limit(50).all())
    sent = failed = 0
    for sub in due:
        try:
            result = send(db, sub)
            sub.last_status = f"Sent {result['rows']} row(s) as {result['file']}"
            sent += 1
        except Exception as exc:  # one bad subscription must not stop the others
            sub.last_status = f"Failed: {str(exc)[:250]}"
            failed += 1
        sub.last_run_at = now
        sub.next_run_at = next_run(sub.frequency, now)
        db.flush()
    return {"due": len(due), "sent": sent, "failed": failed}
