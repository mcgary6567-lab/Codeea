"""Legacy data migration (Module 35 / SRS Section 20).

CSV + XLSX parsing, field mapping with name-similarity suggestions, validation (required / types / duplicates),
transactional per-row import with skip-on-error, and a reconciliation report.
Engine agnostic: pure ORM, no raw SQL.
"""
from __future__ import annotations

import csv
import difflib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Optional

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.config import settings
from app.core.utils import next_code
from app.models.core import User
from app.models.crm import Lead
from app.models.finance import Invoice, LedgerEntry, Payment
from app.models.ops import MigrationJob
from app.models.people import Client, Student, Teacher

MIGRATION_DIR = settings.storage_dir / "uploads" / "migration"

# ----------------------------------------------------------------------------- entity catalogue
# (target column, label, type, required)
ENTITIES: dict[str, dict] = {
    "clients": {
        "label": "Clients / Parents",
        "icon": "users",
        "match_on": "email / phone / client_code",
        "columns": [
            ("client_code", "Legacy client code", "str", False),
            ("full_name", "Full name", "str", True),
            ("email", "Email", "email", False),
            ("phone", "Phone", "phone", False),
            ("whatsapp", "WhatsApp", "phone", False),
            ("country", "Country", "str", False),
            ("city", "City", "str", False),
            ("timezone", "Timezone", "str", False),
            ("currency", "Currency", "str", False),
            ("address", "Address", "str", False),
            ("status", "Status (active/inactive/churned)", "str", False),
            ("joined_at", "Joined date", "date", False),
            ("notes", "Notes", "str", False),
        ],
    },
    "students": {
        "label": "Students",
        "icon": "graduation-cap",
        "match_on": "student_code / client_code + name",
        "columns": [
            ("student_code", "Legacy student code", "str", False),
            ("full_name", "Full name", "str", True),
            ("client_code", "Parent client code", "str", False),
            ("client_email", "Parent email", "email", False),
            ("gender", "Gender", "str", False),
            ("date_of_birth", "Date of birth", "date", False),
            ("level", "Level", "str", False),
            ("status", "Status (active/trial/frozen/cancelled)", "str", False),
            ("join_date", "Join date", "date", False),
            ("timezone", "Timezone", "str", False),
            ("notes", "Notes", "str", False),
        ],
    },
    "leads": {
        "label": "Leads",
        "icon": "funnel",
        "match_on": "email / phone",
        "columns": [
            ("full_name", "Full name", "str", True),
            ("email", "Email", "email", False),
            ("phone", "Phone", "phone", False),
            ("whatsapp", "WhatsApp", "phone", False),
            ("country", "Country", "str", False),
            ("student_name", "Student name", "str", False),
            ("student_age", "Student age", "int", False),
            ("stage", "Stage", "str", False),
            ("notes", "Notes", "str", False),
        ],
    },
    "teachers": {
        "label": "Teachers",
        "icon": "user-check",
        "match_on": "email / teacher_code",
        "columns": [
            ("teacher_code", "Legacy teacher code", "str", False),
            ("full_name", "Full name", "str", True),
            ("email", "Email", "email", False),
            ("phone", "Phone", "phone", False),
            ("gender", "Gender", "str", False),
            ("qualifications", "Qualifications", "str", False),
            ("shift", "Shift", "str", False),
            ("timezone", "Timezone", "str", False),
            ("per_class_rate", "Per class rate", "decimal", False),
            ("status", "Status", "str", False),
        ],
    },
    "invoices": {
        "label": "Invoices",
        "icon": "receipt",
        "match_on": "invoice_number",
        "financial_field": "total",
        "columns": [
            ("invoice_number", "Invoice number", "str", False),
            ("client_code", "Client code", "str", False),
            ("client_email", "Client email", "email", False),
            ("issue_date", "Issue date", "date", True),
            ("due_date", "Due date", "date", False),
            ("currency", "Currency", "str", False),
            ("subtotal", "Subtotal", "decimal", False),
            ("discount", "Discount", "decimal", False),
            ("total", "Total", "decimal", True),
            ("paid_amount", "Paid amount", "decimal", False),
            ("status", "Status (draft/sent/paid/overdue)", "str", False),
            ("remarks", "Remarks", "str", False),
        ],
    },
    "payments": {
        "label": "Payments",
        "icon": "credit-card",
        "match_on": "payment_number / reference",
        "financial_field": "amount",
        "columns": [
            ("payment_number", "Payment number", "str", False),
            ("client_code", "Client code", "str", False),
            ("client_email", "Client email", "email", False),
            ("invoice_number", "Invoice number", "str", False),
            ("amount", "Amount", "decimal", True),
            ("currency", "Currency", "str", False),
            ("method", "Method (bank/card/cash/wise)", "str", False),
            ("reference", "Reference", "str", False),
            ("received_at", "Received at", "date", True),
            ("notes", "Notes", "str", False),
        ],
    },
}

DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"]


def entity_columns(entity: str) -> list[tuple]:
    return ENTITIES.get(entity, {}).get("columns", [])


def required_columns(entity: str) -> list[str]:
    return [c[0] for c in entity_columns(entity) if c[3]]


def ensure_dir() -> Path:
    MIGRATION_DIR.mkdir(parents=True, exist_ok=True)
    return MIGRATION_DIR


# ----------------------------------------------------------------------------- parsing
def save_upload(filename: str, content: bytes) -> Path:
    ensure_dir()
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", Path(filename).name) or "upload.csv"
    stamp = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    path = MIGRATION_DIR / f"{stamp}-{safe}"
    path.write_bytes(content)
    return path


def parse_file(path: str | Path, limit: Optional[int] = None) -> tuple[list[str], list[dict]]:
    """Return (headers, rows-as-dicts). Supports .csv/.txt and .xlsx/.xlsm."""
    p = Path(path)
    if not p.exists():
        return [], []
    suffix = p.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _parse_xlsx(p, limit)
    return _parse_csv(p, limit)


def _clean(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    return str(v).strip()


def _parse_csv(p: Path, limit: Optional[int]) -> tuple[list[str], list[dict]]:
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        text = raw.decode("utf-8", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[dict] = []
    headers: list[str] = []
    for i, row in enumerate(reader):
        if i == 0:
            headers = [_clean(h) or f"column_{n + 1}" for n, h in enumerate(row)]
            continue
        if not any(_clean(c) for c in row):
            continue
        rows.append({headers[n]: _clean(c) for n, c in enumerate(row) if n < len(headers)})
        if limit and len(rows) >= limit:
            break
    return headers, rows


def _parse_xlsx(p: Path, limit: Optional[int]) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook
    wb = load_workbook(filename=str(p), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers: list[str] = []
    rows: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [_clean(h) or f"column_{n + 1}" for n, h in enumerate(row)]
            continue
        if not any(_clean(c) for c in row):
            continue
        rows.append({headers[n]: _clean(c) for n, c in enumerate(row) if n < len(headers)})
        if limit and len(rows) >= limit:
            break
    wb.close()
    return headers, rows


def row_count(path: str | Path) -> int:
    _, rows = parse_file(path)
    return len(rows)


# ----------------------------------------------------------------------------- mapping
def suggest_mapping(entity: str, headers: list[str]) -> dict[str, str]:
    """Auto-suggest target column -> file header using normalised name similarity."""
    def norm(s: str) -> str:
        return re.sub(r"[^a-z0-9]", "", (s or "").lower())

    norm_headers = {norm(h): h for h in headers}
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for target, label, _t, _req in entity_columns(entity):
        candidates = [norm(target), norm(label.split("(")[0]), norm(target.replace("_", " "))]
        hit = None
        for c in candidates:
            if c in norm_headers and norm_headers[c] not in used:
                hit = norm_headers[c]
                break
        if hit is None:
            close = difflib.get_close_matches(norm(target), [k for k in norm_headers if norm_headers[k] not in used], n=1, cutoff=0.78)
            if close:
                hit = norm_headers[close[0]]
        if hit:
            mapping[target] = hit
            used.add(hit)
    return mapping


def map_row(row: dict, mapping: dict) -> dict:
    return {target: (row.get(header) or "").strip() if isinstance(row.get(header), str) else row.get(header)
            for target, header in (mapping or {}).items() if header}


# ----------------------------------------------------------------------------- coercion
def parse_any_date(value: Any) -> Optional[date]:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s[:len(fmt) + 6], fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s).date()
    except ValueError:
        return None


def coerce(value: Any, kind: str) -> tuple[Any, Optional[str]]:
    """Return (converted_value, error_message)."""
    s = "" if value is None else str(value).strip()
    if s == "":
        return None, None
    if kind == "date":
        d = parse_any_date(s)
        return (d, None) if d else (None, f"'{s[:40]}' is not a recognised date")
    if kind == "int":
        try:
            return int(float(s.replace(",", ""))), None
        except (ValueError, TypeError):
            return None, f"'{s[:40]}' is not a whole number"
    if kind == "decimal":
        cleaned = re.sub(r"[^0-9.\-]", "", s.replace(",", ""))
        try:
            return float(Decimal(cleaned)), None
        except (InvalidOperation, ValueError, TypeError):
            return None, f"'{s[:40]}' is not a number"
    if kind == "email":
        if "@" not in s or "." not in s.split("@")[-1]:
            return s, f"'{s[:40]}' is not a valid email address"
        return s.lower(), None
    if kind == "phone":
        digits = re.sub(r"[^0-9+]", "", s)
        if len(re.sub(r"\D", "", digits)) < 7:
            return s, f"'{s[:40]}' is not a valid phone number"
        return digits, None
    return s, None


def _dup_key(entity: str, data: dict) -> Optional[str]:
    if entity == "clients":
        return (data.get("email") or data.get("phone") or data.get("client_code") or "").lower() or None
    if entity == "students":
        return (data.get("student_code") or f"{data.get('client_code', '')}|{data.get('full_name', '')}").lower() or None
    if entity == "leads":
        return (data.get("email") or data.get("phone") or "").lower() or None
    if entity == "teachers":
        return (data.get("email") or data.get("teacher_code") or "").lower() or None
    if entity == "invoices":
        return (data.get("invoice_number") or "").lower() or None
    if entity == "payments":
        return (data.get("payment_number") or data.get("reference") or "").lower() or None
    return None


def _exists_in_db(db: Session, entity: str, data: dict) -> bool:
    if entity == "clients":
        email, phone = data.get("email"), data.get("phone")
        q = db.query(Client.id)
        if email:
            if q.filter(func.lower(Client.email) == str(email).lower()).first():
                return True
        if phone and db.query(Client.id).filter(Client.phone == phone).first():
            return True
        return False
    if entity == "students":
        code = data.get("student_code")
        if code and db.query(Student.id).filter(Student.student_code == code).first():
            return True
        name, ccode = data.get("full_name"), data.get("client_code")
        if name and ccode:
            client = db.query(Client).filter(Client.client_code == ccode).first()
            if client and db.query(Student.id).filter(Student.client_id == client.id, Student.full_name == name).first():
                return True
        return False
    if entity == "leads":
        email, phone = data.get("email"), data.get("phone")
        if email and db.query(Lead.id).filter(func.lower(Lead.email) == str(email).lower()).first():
            return True
        if phone and db.query(Lead.id).filter(Lead.phone == phone).first():
            return True
        return False
    if entity == "teachers":
        email = data.get("email")
        if email and db.query(Teacher).join(User, Teacher.user_id == User.id).filter(func.lower(User.email) == str(email).lower()).first():
            return True
        return False
    if entity == "invoices":
        num = data.get("invoice_number")
        return bool(num and db.query(Invoice.id).filter(Invoice.invoice_number == num).first())
    if entity == "payments":
        num = data.get("payment_number")
        if num and db.query(Payment.id).filter(Payment.payment_number == num).first():
            return True
        ref = data.get("reference")
        return bool(ref and db.query(Payment.id).filter(Payment.reference == ref).first())
    return False


# ----------------------------------------------------------------------------- validation
def validate_job(db: Session, job: MigrationJob) -> dict:
    """Check required fields, types, in-file duplicates and database duplicates. Stores errors on the job."""
    headers, rows = parse_file(job.file_path or "")
    mapping = job.field_mapping or {}
    cols = {c[0]: c for c in entity_columns(job.entity)}
    errors: list[dict] = []
    seen: dict[str, int] = {}
    duplicates = 0
    for idx, row in enumerate(rows, start=2):  # row 1 is the header
        data = map_row(row, mapping)
        clean: dict[str, Any] = {}
        for target, raw in data.items():
            spec = cols.get(target)
            if not spec:
                continue
            value, err = coerce(raw, spec[2])
            clean[target] = value
            if err:
                errors.append({"row": idx, "field": target, "message": err, "severity": "error"})
        for req in required_columns(job.entity):
            if clean.get(req) in (None, ""):
                errors.append({"row": idx, "field": req, "message": f"Required field '{req}' is empty or unmapped", "severity": "error"})
        key = _dup_key(job.entity, {k: v for k, v in clean.items() if v is not None})
        if key:
            if key in seen:
                duplicates += 1
                errors.append({"row": idx, "field": "duplicate", "message": f"Duplicate of row {seen[key]} in this file (key {key[:60]})", "severity": "warning"})
            else:
                seen[key] = idx
        if _exists_in_db(db, job.entity, {k: v for k, v in clean.items() if v is not None}):
            duplicates += 1
            errors.append({"row": idx, "field": "duplicate", "message": "A matching record already exists in the database - row will be skipped", "severity": "warning"})
    job.records_total = len(rows)
    job.duplicates_found = duplicates
    job.errors = errors[:500]
    job.status = "validated"
    hard = sum(1 for e in errors if e["severity"] == "error")
    return {"rows": len(rows), "errors": hard, "warnings": len(errors) - hard, "duplicates": duplicates,
            "unmapped_required": [c for c in required_columns(job.entity) if c not in mapping]}


# ----------------------------------------------------------------------------- import
def _client_for(db: Session, data: dict) -> Optional[Client]:
    code = data.get("client_code")
    if code:
        c = db.query(Client).filter(Client.client_code == code).first()
        if c:
            return c
    email = data.get("client_email") or data.get("email")
    if email:
        c = db.query(Client).filter(func.lower(Client.email) == str(email).lower()).first()
        if c:
            return c
    return None


def _ledger(db: Session, client_id: int, entry_date: date, entry_type: str, description: str,
            debit: float = 0.0, credit: float = 0.0, currency: str = "PKR", ref_type: str = "", ref_id: int = 0) -> LedgerEntry:
    last = (db.query(LedgerEntry).filter(LedgerEntry.client_id == client_id)
            .order_by(LedgerEntry.entry_date.desc(), LedgerEntry.id.desc()).first())
    balance = float(last.balance_after or 0) if last else 0.0
    balance = balance + float(debit or 0) - float(credit or 0)
    e = LedgerEntry(client_id=client_id, entry_date=entry_date, entry_type=entry_type, description=description,
                    debit=debit, credit=credit, currency=currency, balance_after=balance,
                    reference_type=ref_type or None, reference_id=ref_id or None)
    db.add(e)
    return e


def _import_row(db: Session, entity: str, clean: dict, actor: Optional[User]) -> str:
    """Insert one record. Returns a short label. Raises ValueError to skip the row."""
    if entity == "clients":
        obj = Client(
            client_code=clean.get("client_code") or next_code(db, Client, "client_code", "C-"),
            full_name=clean["full_name"], email=clean.get("email"), phone=clean.get("phone"),
            whatsapp=clean.get("whatsapp") or clean.get("phone"), country=clean.get("country") or "Pakistan",
            city=clean.get("city"), timezone=clean.get("timezone") or settings.DEFAULT_TIMEZONE,
            currency=clean.get("currency") or settings.BASE_CURRENCY, address=clean.get("address"),
            status=clean.get("status") or "active", notes=clean.get("notes"), source="migration",
            joined_at=clean.get("joined_at") or date.today())
        if db.query(Client).filter(Client.client_code == obj.client_code).first():
            raise ValueError(f"client_code {obj.client_code} already exists")
        db.add(obj)
        db.flush()
        return obj.client_code
    if entity == "students":
        client = _client_for(db, clean)
        if not client:
            raise ValueError("no matching client (client_code / client_email)")
        obj = Student(
            student_code=clean.get("student_code") or next_code(db, Student, "student_code", "S-"),
            client_id=client.id, full_name=clean["full_name"], gender=clean.get("gender") or "male",
            date_of_birth=clean.get("date_of_birth"), level=clean.get("level"),
            status=clean.get("status") or "active", timezone=clean.get("timezone") or client.timezone,
            join_date=clean.get("join_date") or date.today(), notes=clean.get("notes"))
        if db.query(Student).filter(Student.student_code == obj.student_code).first():
            raise ValueError(f"student_code {obj.student_code} already exists")
        db.add(obj)
        db.flush()
        return obj.student_code
    if entity == "leads":
        obj = Lead(lead_code=next_code(db, Lead, "lead_code", "L-"), full_name=clean["full_name"],
                   email=clean.get("email"), phone=clean.get("phone"), whatsapp=clean.get("whatsapp"),
                   country=clean.get("country"), student_name=clean.get("student_name"),
                   student_age=clean.get("student_age"), stage=clean.get("stage") or "new", notes=clean.get("notes"))
        db.add(obj)
        db.flush()
        return obj.lead_code
    if entity == "teachers":
        obj = Teacher(teacher_code=clean.get("teacher_code") or next_code(db, Teacher, "teacher_code", "T-"),
                      full_name=clean["full_name"], gender=clean.get("gender") or "male",
                      qualifications=clean.get("qualifications"), shift=clean.get("shift") or "morning",
                      timezone=clean.get("timezone") or settings.DEFAULT_TIMEZONE,
                      per_class_rate=clean.get("per_class_rate") or 0, status=clean.get("status") or "active")
        if db.query(Teacher).filter(Teacher.teacher_code == obj.teacher_code).first():
            raise ValueError(f"teacher_code {obj.teacher_code} already exists")
        db.add(obj)
        db.flush()
        return obj.teacher_code
    if entity == "invoices":
        client = _client_for(db, clean)
        if not client:
            raise ValueError("no matching client (client_code / client_email)")
        total = float(clean.get("total") or 0)
        issue = clean.get("issue_date") or date.today()
        obj = Invoice(invoice_number=clean.get("invoice_number") or next_code(db, Invoice, "invoice_number", "INV-", width=6),
                      client_id=client.id, issue_date=issue, due_date=clean.get("due_date") or issue,
                      currency=clean.get("currency") or client.currency,
                      subtotal=clean.get("subtotal") if clean.get("subtotal") is not None else total,
                      discount=clean.get("discount") or 0, total=total,
                      paid_amount=clean.get("paid_amount") or 0, total_in_base=total,
                      status=clean.get("status") or ("paid" if float(clean.get("paid_amount") or 0) >= total > 0 else "sent"),
                      remarks=clean.get("remarks") or "Imported from legacy system")
        if db.query(Invoice).filter(Invoice.invoice_number == obj.invoice_number).first():
            raise ValueError(f"invoice_number {obj.invoice_number} already exists")
        db.add(obj)
        db.flush()
        _ledger(db, client.id, issue, "invoice", f"Invoice {obj.invoice_number} (migrated)", debit=total,
                currency=obj.currency, ref_type="Invoice", ref_id=obj.id)
        return obj.invoice_number
    if entity == "payments":
        client = _client_for(db, clean)
        if not client:
            raise ValueError("no matching client (client_code / client_email)")
        amount = float(clean.get("amount") or 0)
        received = clean.get("received_at") or date.today()
        invoice = None
        if clean.get("invoice_number"):
            invoice = db.query(Invoice).filter(Invoice.invoice_number == clean["invoice_number"]).first()
        obj = Payment(payment_number=clean.get("payment_number") or next_code(db, Payment, "payment_number", "PAY-", width=6),
                      client_id=client.id, invoice_id=invoice.id if invoice else None, amount=amount,
                      currency=clean.get("currency") or client.currency, amount_in_base=amount,
                      method=clean.get("method") or "bank", reference=clean.get("reference"),
                      status="succeeded", received_at=datetime.combine(received, datetime.min.time()),
                      notes=clean.get("notes") or "Imported from legacy system")
        if db.query(Payment).filter(Payment.payment_number == obj.payment_number).first():
            raise ValueError(f"payment_number {obj.payment_number} already exists")
        db.add(obj)
        db.flush()
        if invoice:
            invoice.paid_amount = float(invoice.paid_amount or 0) + amount
            if float(invoice.paid_amount) >= float(invoice.total or 0):
                invoice.status = "paid"
        _ledger(db, client.id, received, "payment", f"Payment {obj.payment_number} (migrated)", credit=amount,
                currency=obj.currency, ref_type="Payment", ref_id=obj.id)
        return obj.payment_number
    raise ValueError(f"unsupported entity {entity}")


def import_job(db: Session, job: MigrationJob, actor: Optional[User]) -> dict:
    """Per-row transactional import with skip-on-error. Returns a summary dict."""
    headers, rows = parse_file(job.file_path or "")
    mapping = job.field_mapping or {}
    cols = {c[0]: c for c in entity_columns(job.entity)}
    imported, skipped, duplicates = 0, 0, 0
    errors: list[dict] = list(job.errors or [])[:0]
    seen: set[str] = set()
    financial_total = 0.0
    fin_field = ENTITIES.get(job.entity, {}).get("financial_field")
    for idx, row in enumerate(rows, start=2):
        data = map_row(row, mapping)
        clean: dict[str, Any] = {}
        row_error: Optional[str] = None
        for target, raw in data.items():
            spec = cols.get(target)
            if not spec:
                continue
            value, err = coerce(raw, spec[2])
            if err and spec[3]:
                row_error = err
            if value not in (None, ""):
                clean[target] = value
        for req in required_columns(job.entity):
            if clean.get(req) in (None, ""):
                row_error = f"Required field '{req}' is empty"
        key = _dup_key(job.entity, clean)
        if not row_error and key and key in seen:
            row_error = "Duplicate row within the file"
            duplicates += 1
        if not row_error and _exists_in_db(db, job.entity, clean):
            row_error = "Record already exists in the database"
            duplicates += 1
        if row_error:
            skipped += 1
            errors.append({"row": idx, "field": "-", "message": row_error, "severity": "warning"})
            continue
        savepoint = db.begin_nested()
        try:
            label = _import_row(db, job.entity, clean, actor)
            savepoint.commit()
            imported += 1
            if key:
                seen.add(key)
            if fin_field and clean.get(fin_field) is not None:
                financial_total += float(clean[fin_field])
        except Exception as exc:
            savepoint.rollback()
            skipped += 1
            errors.append({"row": idx, "field": "-", "message": f"{type(exc).__name__}: {str(exc)[:180]}", "severity": "error"})
    job.records_total = len(rows)
    job.records_imported = imported
    job.records_skipped = skipped
    job.duplicates_found = duplicates
    job.errors = errors[:500]
    job.status = "imported" if imported else "failed"
    job.completed_at = datetime.utcnow()
    return {"imported": imported, "skipped": skipped, "duplicates": duplicates, "rows": len(rows),
            "financial_total": round(financial_total, 2)}


# ----------------------------------------------------------------------------- reconciliation
def reconciliation(db: Session, job: MigrationJob) -> dict:
    headers, rows = parse_file(job.file_path or "")
    mapping = job.field_mapping or {}
    fin_field = ENTITIES.get(job.entity, {}).get("financial_field")
    file_total = 0.0
    if fin_field and mapping.get(fin_field):
        for row in rows:
            value, err = coerce(row.get(mapping[fin_field]), "decimal")
            if value:
                file_total += float(value)
    imported_total = 0.0
    if job.entity == "invoices":
        imported_total = float(db.query(func.coalesce(func.sum(Invoice.total), 0))
                               .filter(Invoice.remarks.like("%legacy%")).scalar() or 0)
    elif job.entity == "payments":
        imported_total = float(db.query(func.coalesce(func.sum(Payment.amount), 0))
                               .filter(Payment.notes.like("%legacy%")).scalar() or 0)
    hard = [e for e in (job.errors or []) if e.get("severity") == "error"]
    warn = [e for e in (job.errors or []) if e.get("severity") != "error"]
    return {
        "file_rows": len(rows), "headers": headers, "mapped_fields": len([v for v in mapping.values() if v]),
        "records_total": job.records_total, "records_imported": job.records_imported,
        "records_skipped": job.records_skipped, "duplicates": job.duplicates_found,
        "financial_field": fin_field, "file_total": round(file_total, 2), "imported_total": round(imported_total, 2),
        "variance": round(imported_total - file_total, 2) if fin_field else 0.0,
        "hard_errors": len(hard), "warnings": len(warn),
        "success_pct": round(100.0 * job.records_imported / job.records_total, 1) if job.records_total else 0.0,
    }


# ----------------------------------------------------------------------------- CSV helpers
def template_csv(entity: str) -> str:
    cols = entity_columns(entity)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c[0] for c in cols])
    sample = {
        "clients": ["LEG-1001", "Ahmed Khan", "ahmed.khan@example.com", "+447700900123", "+447700900123", "United Kingdom",
                    "Birmingham", "Europe/London", "GBP", "12 Rose Lane", "active", "2024-03-14", "Migrated from legacy ERP"],
        "students": ["LEG-S-2001", "Zainab Ahmed", "LEG-1001", "ahmed.khan@example.com", "female", "2014-06-02",
                     "Nazira 2", "active", "2024-03-20", "Europe/London", "Reads fluently"],
        "leads": ["Bilal Yusuf", "bilal@example.com", "+923001234567", "+923001234567", "Pakistan", "Hamza", "9", "new", "Facebook enquiry"],
        "teachers": ["LEG-T-11", "Qari Abdul Rehman", "abdul.rehman@example.com", "+923011234567", "male",
                     "Hafiz-e-Quran, Tajweed Ijazah", "morning", "Asia/Karachi", "350", "active"],
        "invoices": ["LEG-INV-5001", "LEG-1001", "ahmed.khan@example.com", "2025-08-01", "2025-08-07", "GBP",
                     "45.00", "0.00", "45.00", "45.00", "paid", "August tuition"],
        "payments": ["LEG-PAY-9001", "LEG-1001", "ahmed.khan@example.com", "LEG-INV-5001", "45.00", "GBP",
                     "bank", "TRX-88213", "2025-08-03", "Bank transfer"],
    }.get(entity)
    if sample:
        w.writerow(sample[:len(cols)])
    return buf.getvalue()


def errors_csv(job: MigrationJob) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["row", "field", "severity", "message"])
    for e in (job.errors or []):
        w.writerow([e.get("row"), e.get("field"), e.get("severity"), e.get("message")])
    return buf.getvalue()


# ----------------------------------------------------------------------------- Section 20 cutover checklist
DEFAULT_CHECKLIST = [
    {"key": "inventory", "phase": "1. Inventory", "title": "Inventory every legacy system",
     "detail": "List the legacy ERP, spreadsheets, WhatsApp exports, Google Sheets, accounting files and paper records. Record owner, record count, format and refresh frequency.", "done": False},
    {"key": "field_map", "phase": "1. Inventory", "title": "Map legacy fields to OQC entities",
     "detail": "For each source file, map its columns onto clients / students / leads / teachers / invoices / payments using the wizard's mapping step. Note fields that have no home.", "done": False},
    {"key": "cleanse", "phase": "2. Cleansing", "title": "De-duplicate and cleanse at source",
     "detail": "Fix names, normalise phone numbers to E.164, resolve duplicate families, and drop test rows before exporting.", "done": False},
    {"key": "templates", "phase": "2. Cleansing", "title": "Export into the OQC CSV templates",
     "detail": "Download the per-entity template from this page and re-shape the legacy export into it. Keep one file per entity.", "done": False},
    {"key": "dry_run", "phase": "3. Dry run", "title": "Validate on a staging database",
     "detail": "Run upload -> map -> validate against staging. Fix every hard error; review each duplicate warning.", "done": False},
    {"key": "order", "phase": "3. Dry run", "title": "Confirm the import order",
     "detail": "Clients, then students, then teachers, then invoices, then payments, then leads. Children need their parents to exist.", "done": False},
    {"key": "backup", "phase": "4. Cutover", "title": "Take a full backup immediately before cutover",
     "detail": "Create a backup in /admin/backups and run a restore test on it. This is the rollback point.", "done": False},
    {"key": "freeze", "phase": "4. Cutover", "title": "Freeze the legacy system",
     "detail": "Announce a write freeze on the legacy ERP so no new records appear mid-migration.", "done": False},
    {"key": "import", "phase": "4. Cutover", "title": "Run the production import",
     "detail": "Import each entity in order, checking the reconciliation report after every run.", "done": False},
    {"key": "reconcile", "phase": "5. Verify", "title": "Reconcile counts and financial totals",
     "detail": "Row counts, client balances and the imported financial total must match the legacy figures within the agreed tolerance.", "done": False},
    {"key": "spot_check", "phase": "5. Verify", "title": "Spot-check 20 records end to end",
     "detail": "Pick 20 families across countries and check profile, students, subscription, invoices, ledger balance and portal login.", "done": False},
    {"key": "parallel", "phase": "6. Stabilise", "title": "Run parallel for one billing cycle",
     "detail": "Keep the legacy system read-only for one month and compare invoices and payments before decommissioning.", "done": False},
    {"key": "train", "phase": "6. Stabilise", "title": "Train staff and publish the new SOPs",
     "detail": "Every department signs off that their daily workflow works in the OS.", "done": False},
    {"key": "decommission", "phase": "6. Stabilise", "title": "Archive and decommission the legacy system",
     "detail": "Take a final legacy export, store it with the off-site backups, then revoke access.", "done": False},
]


def checklist(db: Session) -> list[dict]:
    from app.services.system import get_setting
    stored = get_setting(db, "migration_checklist")
    state = {}
    if isinstance(stored, dict):
        state = stored.get("items", stored)
    out = []
    for item in DEFAULT_CHECKLIST:
        row = dict(item)
        if isinstance(state, dict):
            row["done"] = bool(state.get(item["key"], item["done"]))
        out.append(row)
    return out
